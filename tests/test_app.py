import copy
import tempfile
import unittest
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock
from openpyxl import load_workbook
from aiogram import Bot, Dispatcher, types
from waiterbot.app import TrainingBot, chunks, export_results, Guard
from waiterbot.catalog import Catalog, units
from waiterbot.storage import Store


class AppTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.store = Store(Path(temporary.name) / "bot.sqlite3")
        self.store.initialize()
        self.store.set_role(101, "waiter")
        self.store.upsert_user(101, "=HYPERLINK(test)", "employee")
        self.catalog = Catalog(Path(__file__).resolve().parents[1])
        self.catalog.reload()
        self.app = TrainingBot(self.store, self.catalog, admin_id=999)
        self.message = SimpleNamespace(answer=AsyncMock(), edit_text=AsyncMock(), answer_document=AsyncMock())

    async def press(self, data, user=101):
        await self.app.callback(SimpleNamespace(message=self.message, from_user=SimpleNamespace(id=user), data=data))

    def visible(self):
        return self.message.edit_text.call_args.args[0]

    async def test_complete_train_restart_repeat_and_report(self):
        question = self.catalog.questions["waiter"][0]
        session = self.app.training.start(101, "waiter", question["category"], [question], 1)
        wrong = question["correct"] % 4 + 1
        await self.press(f"answer:{session['id']}:0:{wrong}")
        self.assertIn("Правильный ответ:", self.visible())
        # The final result is saved before the summary button is clicked.
        self.assertEqual(len(self.store.read_results(mode="test")), 1)
        self.app = TrainingBot(Store(self.store.db_path), self.catalog, 999)
        await self.press("menu:resume")
        self.assertIn("Правильный ответ:", self.visible())
        self.assertEqual(len(self.store.read_results()), 1)
        await self.press(f"next:{session['id']}:1")
        self.assertIn("Результат сохранён", self.visible())
        await self.press("menu:practice")
        practice = self.store.get_session(101)
        self.assertEqual(practice["mode"], "practice")
        await self.press(f"answer:{practice['id']}:0:{question['correct']}")
        await self.press(f"next:{practice['id']}:1")
        self.assertEqual(self.store.get_mistakes(101), [])
        self.assertEqual(len(self.store.read_results(mode="test")), 1)
        data = export_results(self.store.read_results())
        book = load_workbook(BytesIO(data))
        try:
            self.assertEqual(book.active.max_row, 3)
            self.assertEqual(book.active["C2"].data_type, "s")
            self.assertEqual(book.active["C2"].value, "=HYPERLINK(test)")
        finally:
            book.close()

    async def test_all_admin_routes_reject_forged_non_admin_callbacks(self):
        for data in ("menu:admin", "admin:download", "admin:reload", "admin:stats",
                     "qcount_role:waiter", "qcount_set:waiter:30"):
            with self.subTest(data=data), self.assertRaisesRegex(ValueError, "администратора"):
                await self.press(data)
        self.message.answer_document.assert_not_awaited()
        self.assertEqual(self.store.get_question_count("waiter"), 10)

    async def test_long_categories_use_short_stable_callbacks(self):
        self.catalog.questions["waiter"] = [dict(self.catalog.questions["waiter"][0], category="Длинный раздел " * 10)]
        await self.press("menu:test")
        markup = self.message.edit_text.call_args.kwargs["reply_markup"]
        for row in markup.inline_keyboard:
            for button in row:
                self.assertLessEqual(len(button.callback_data.encode()), 64)
        await self.press(markup.inline_keyboard[0][0].callback_data)
        self.assertIsNotNone(self.store.get_session(101))

    async def test_old_role_and_old_cancel_buttons_preserve_active_test(self):
        question = self.catalog.questions["waiter"][0]
        session = self.app.training.start(101, "waiter", question["category"], [question], 1)
        await self.press("set_role:cook")
        self.assertEqual(self.store.get_role(101), "waiter")
        self.assertEqual(self.store.get_session(101), session)
        with self.assertRaises(ValueError):
            await self.press("cancel_yes:obsolete")
        await self.press("cancel:" + session["id"])
        self.assertIsNotNone(self.store.get_session(101))
        await self.press("cancel_yes:" + session["id"])
        self.assertIsNone(self.store.get_session(101))

    async def test_real_aiogram_dispatcher_routes_start_with_mock_transport(self):
        bot = Bot(token="123456:TEST_ONLY_FAKE_TOKEN_abcdefghijklmnopqrstuvwxyz")
        bot.session = AsyncMock()
        self.addAsyncCleanup(bot.session.close)
        dispatcher = Dispatcher()
        dispatcher.include_router(self.app.router)
        update = types.Update(update_id=1, message=types.Message(
            message_id=1, date=datetime.now(timezone.utc), chat=types.Chat(id=101, type="private"),
            from_user=types.User(id=101, is_bot=False, first_name="Сотрудник"), text="/start"))
        await dispatcher.feed_update(bot, update)
        methods = [call.args[1] for call in bot.session.call_args_list]
        self.assertEqual(len(methods), 2)
        self.assertIn("Обучение команды", methods[-1].text)

    async def test_group_updates_never_expose_private_results(self):
        guard = Guard(self.app)
        handler = AsyncMock()
        message = types.Message(message_id=2, date=datetime.now(timezone.utc),
            chat=types.Chat(id=-1001, type="supergroup"), from_user=types.User(id=101, is_bot=False, first_name="Сотрудник"),
            text="/results")
        await guard(handler, message, {})
        handler.assert_not_awaited()

    def test_telegram_chunks_preserve_text_and_utf16_limits(self):
        text = "😀Текст\n" * 1000
        parts = list(chunks(text))
        self.assertEqual("".join(parts), text)
        self.assertTrue(all(0 < units(part) <= 3500 for part in parts))
