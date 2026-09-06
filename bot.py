import asyncio
import os
import json
import random
from datetime import datetime
from collections import defaultdict

import openpyxl
from openpyxl import Workbook
from dotenv import load_dotenv

from aiogram import Bot, Dispatcher, types, F
from aiogram.filters import CommandStart, Command
from aiogram.types import (
    InlineKeyboardMarkup, InlineKeyboardButton, FSInputFile,
    ReplyKeyboardMarkup, KeyboardButton,
)

# =====================================================================
# НАСТРОЙКИ
# =====================================================================
load_dotenv()
TOKEN = os.getenv("BOT_TOKEN")
ADMIN_ID = int(os.getenv("ADMIN_ID", "0"))

bot = Bot(token=TOKEN)
dp = Dispatcher()

# ---------------------------------------------------------------------
# СПИСОК ПРОФЕССИЙ И ИХ ФАЙЛОВ С ВОПРОСАМИ.
# Чтобы добавить новую профессию: положите файл вопросов рядом с bot.py
# в том же формате (7 колонок) и добавьте одну строку сюда.
# ---------------------------------------------------------------------
ROLES = {
    "waiter":    {"label": "🍽 Официант",  "file": "questions_waiter.xlsx"},
    "pizzaiolo": {"label": "🍕 Пиццайоло", "file": "questions_pizzaiolo.xlsx"},
    "cook":      {"label": "👨‍🍳 Повар",    "file": "questions_cook.xlsx"},
}

ROLES_FILE = "user_roles.json"
SETTINGS_FILE = "settings.json"
RESULTS_FILE = "results.xlsx"
DEFAULT_QUESTIONS_PER_TEST = 10
QUESTION_COUNT_OPTIONS = [5, 10, 15, 20, 25, 30]


# =====================================================================
# ЗАГРУЗКА ВОПРОСОВ (отдельно на каждую профессию, с кэшем)
# =====================================================================
_questions_cache = {}


def load_questions_for_role(role_key):
    if role_key in _questions_cache and _questions_cache[role_key]:
        return _questions_cache[role_key]

    path = ROLES[role_key]["file"]
    questions = []
    if os.path.exists(path):
        wb = openpyxl.load_workbook(path)
        ws = wb.active
        for row in ws.iter_rows(min_row=2, values_only=True):
            category, question, opt1, opt2, opt3, opt4, correct = row
            if not question:
                continue
            questions.append({
                "category": category,
                "question": question,
                "options": [opt1, opt2, opt3, opt4],
                "correct": int(correct),
            })
    if questions:
        _questions_cache[role_key] = questions
    return questions


def categories_for_role(role_key):
    return sorted(set(q["category"] for q in load_questions_for_role(role_key)))


user_sessions = {}  # активные прохождения теста


# =====================================================================
# ХРАНЕНИЕ ВЫБРАННОЙ ПРОФЕССИИ (простой json-файл user_id -> role_key)
# =====================================================================
def load_user_roles():
    if os.path.exists(ROLES_FILE):
        with open(ROLES_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    return {}


def save_user_role(user_id, role_key):
    roles = load_user_roles()
    roles[str(user_id)] = role_key
    with open(ROLES_FILE, "w", encoding="utf-8") as f:
        json.dump(roles, f, ensure_ascii=False, indent=2)


def get_user_role(user_id):
    roles = load_user_roles()
    return roles.get(str(user_id))


# =====================================================================
# НАСТРОЙКИ: КОЛИЧЕСТВО ВОПРОСОВ НА ТЕСТ (отдельно на каждую профессию)
# =====================================================================
def load_settings():
    if os.path.exists(SETTINGS_FILE):
        with open(SETTINGS_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    return {}


def get_questions_per_test(role_key):
    settings = load_settings()
    return settings.get(role_key, DEFAULT_QUESTIONS_PER_TEST)


def save_questions_per_test(role_key, value):
    settings = load_settings()
    settings[role_key] = value
    with open(SETTINGS_FILE, "w", encoding="utf-8") as f:
        json.dump(settings, f, ensure_ascii=False, indent=2)


# =====================================================================
# РЕЗУЛЬТАТЫ
# =====================================================================
RESULTS_HEADERS = ["Дата и время", "User ID", "Имя", "Username", "Профессия", "Раздел", "Правильных", "Всего", "Процент"]


def ensure_results_file():
    if not os.path.exists(RESULTS_FILE):
        wb = Workbook()
        ws = wb.active
        ws.append(RESULTS_HEADERS)
        wb.save(RESULTS_FILE)


def save_result(user, role_key, category, score, total):
    ensure_results_file()
    wb = openpyxl.load_workbook(RESULTS_FILE)
    ws = wb.active
    percent = round(score / total * 100)
    role_label = ROLES.get(role_key, {}).get("label", role_key)
    ws.append([
        datetime.now().strftime("%d.%m.%Y %H:%M"),
        user.id,
        user.full_name,
        f"@{user.username}" if user.username else "-",
        role_label,
        category,
        score,
        total,
        percent,
    ])
    wb.save(RESULTS_FILE)


def read_all_results():
    ensure_results_file()
    wb = openpyxl.load_workbook(RESULTS_FILE)
    ws = wb.active
    return list(ws.iter_rows(min_row=2, values_only=True))
    # (дата, user_id, имя, username, профессия, раздел, правильных, всего, процент)


# =====================================================================
# КЛАВИАТУРЫ
# =====================================================================
def role_selection_keyboard():
    buttons = [[InlineKeyboardButton(text=r["label"], callback_data=f"set_role:{key}")] for key, r in ROLES.items()]
    return InlineKeyboardMarkup(inline_keyboard=buttons)


def main_menu_keyboard(is_admin_user: bool):
    buttons = [
        [InlineKeyboardButton(text="📝 Пройти тест", callback_data="menu:test")],
        [InlineKeyboardButton(text="📊 Мои результаты", callback_data="menu:my_results")],
        [InlineKeyboardButton(text="🔁 Сменить профессию", callback_data="menu:change_role")],
    ]
    if is_admin_user:
        buttons.append([InlineKeyboardButton(text="⚙️ Админ-панель", callback_data="menu:admin")])
    return InlineKeyboardMarkup(inline_keyboard=buttons)


def categories_keyboard(role_key):
    cats = categories_for_role(role_key)
    buttons = [[InlineKeyboardButton(text=cat, callback_data=f"start_test:{cat}")] for cat in cats]
    buttons.append([InlineKeyboardButton(text="⬅️ Назад", callback_data="menu:home")])
    return InlineKeyboardMarkup(inline_keyboard=buttons)


def back_to_menu_keyboard():
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="⬅️ В главное меню", callback_data="menu:home")]
    ])


def admin_menu_keyboard():
    buttons = [
        [InlineKeyboardButton(text="📈 Общая статистика", callback_data="admin:stats")],
        [InlineKeyboardButton(text="📉 Слабые результаты (<60%)", callback_data="admin:weak")],
        [InlineKeyboardButton(text="🔢 Кол-во вопросов", callback_data="admin:qcount")],
        [InlineKeyboardButton(text="📥 Скачать файл результатов", callback_data="admin:download")],
        [InlineKeyboardButton(text="⬅️ В главное меню", callback_data="menu:home")],
    ]
    return InlineKeyboardMarkup(inline_keyboard=buttons)


def qcount_roles_keyboard():
    buttons = []
    for key, r in ROLES.items():
        current = get_questions_per_test(key)
        current_label = "все" if current == "all" else str(current)
        buttons.append([InlineKeyboardButton(
            text=f'{r["label"]}: {current_label}', callback_data=f"qcount_role:{key}"
        )])
    buttons.append([InlineKeyboardButton(text="⬅️ Назад", callback_data="menu:admin")])
    return InlineKeyboardMarkup(inline_keyboard=buttons)


def qcount_values_keyboard(role_key):
    buttons = []
    row = []
    for val in QUESTION_COUNT_OPTIONS:
        row.append(InlineKeyboardButton(text=str(val), callback_data=f"qcount_set:{role_key}:{val}"))
        if len(row) == 3:
            buttons.append(row)
            row = []
    if row:
        buttons.append(row)
    buttons.append([InlineKeyboardButton(text="Все вопросы раздела", callback_data=f"qcount_set:{role_key}:all")])
    buttons.append([InlineKeyboardButton(text="⬅️ Назад", callback_data="admin:qcount")])
    return InlineKeyboardMarkup(inline_keyboard=buttons)


def persistent_keyboard(is_admin_user: bool):
    buttons = [
        [KeyboardButton(text="📝 Пройти тест"), KeyboardButton(text="📊 Мои результаты")],
    ]
    if is_admin_user:
        buttons.append([KeyboardButton(text="⚙️ Админ-панель")])
    return ReplyKeyboardMarkup(keyboard=buttons, resize_keyboard=True, is_persistent=True)


def is_admin(user_id):
    return ADMIN_ID != 0 and user_id == ADMIN_ID


# =====================================================================
# ГЛАВНОЕ МЕНЮ / ВЫБОР ПРОФЕССИИ
# =====================================================================
async def show_main_menu(chat_id, message_id=None, user_id=None):
    is_adm = is_admin(user_id)
    role_key = get_user_role(user_id)
    role_label = ROLES.get(role_key, {}).get("label", "не выбрана")
    text = (
        "👋 Привет! Я бот для тестирования персонала по меню и алкоголю.\n\n"
        f"Твоя профессия: {role_label}\n\n"
        "Выбери, что хочешь сделать:"
    )
    kb = main_menu_keyboard(is_adm)
    if message_id:
        await bot.edit_message_text(text, chat_id=chat_id, message_id=message_id, reply_markup=kb)
    else:
        await bot.send_message(chat_id, text, reply_markup=kb)


@dp.message(CommandStart())
async def start_handler(message: types.Message):
    role_key = get_user_role(message.from_user.id)
    is_adm = is_admin(message.from_user.id)
    await message.answer("Меню открыто снизу 👇", reply_markup=persistent_keyboard(is_adm))

    if not role_key:
        await message.answer("Сначала укажи, кто ты:", reply_markup=role_selection_keyboard())
    else:
        await show_main_menu(message.chat.id, user_id=message.from_user.id)


@dp.callback_query(F.data.startswith("set_role:"))
async def set_role(callback: types.CallbackQuery):
    role_key = callback.data.split(":", 1)[1]
    save_user_role(callback.from_user.id, role_key)
    await callback.message.edit_text(f"Профессия сохранена: {ROLES[role_key]['label']}")
    await show_main_menu(callback.message.chat.id, user_id=callback.from_user.id)
    await callback.answer()


@dp.callback_query(F.data == "menu:change_role")
async def change_role(callback: types.CallbackQuery):
    await callback.message.edit_text("Кто ты?", reply_markup=role_selection_keyboard())
    await callback.answer()


@dp.callback_query(F.data == "menu:home")
async def menu_home(callback: types.CallbackQuery):
    await show_main_menu(callback.message.chat.id, message_id=callback.message.message_id, user_id=callback.from_user.id)
    await callback.answer()


@dp.callback_query(F.data == "menu:test")
async def menu_test(callback: types.CallbackQuery):
    role_key = get_user_role(callback.from_user.id)
    if not role_key:
        await callback.message.edit_text("Сначала укажи, кто ты:", reply_markup=role_selection_keyboard())
        await callback.answer()
        return
    await callback.message.edit_text("Выбери раздел для теста:", reply_markup=categories_keyboard(role_key))
    await callback.answer()


# ---------- кнопки постоянного меню внизу экрана ----------
@dp.message(F.text == "📝 Пройти тест")
async def kb_test(message: types.Message):
    role_key = get_user_role(message.from_user.id)
    if not role_key:
        await message.answer("Сначала укажи, кто ты:", reply_markup=role_selection_keyboard())
        return
    await message.answer("Выбери раздел для теста:", reply_markup=categories_keyboard(role_key))


@dp.message(F.text == "📊 Мои результаты")
async def kb_my_results(message: types.Message):
    await send_my_results(message.chat.id, message.from_user.id)


@dp.message(F.text == "⚙️ Админ-панель")
async def kb_admin(message: types.Message):
    if not is_admin(message.from_user.id):
        return
    await message.answer("⚙️ Админ-панель", reply_markup=admin_menu_keyboard())


@dp.message(Command("test"))
async def test_command(message: types.Message):
    role_key = get_user_role(message.from_user.id)
    if not role_key:
        await message.answer("Сначала укажи, кто ты:", reply_markup=role_selection_keyboard())
        return
    await message.answer("Выбери раздел для теста:", reply_markup=categories_keyboard(role_key))


# =====================================================================
# МОИ РЕЗУЛЬТАТЫ
# =====================================================================
async def send_my_results(chat_id, user_id, message_id=None):
    all_results = read_all_results()
    my_results = [r for r in all_results if r[1] == user_id][-10:]

    if not my_results:
        text = "У тебя пока нет пройденных тестов.\nНажми «Пройти тест», чтобы начать!"
    else:
        lines = ["📊 Твои последние результаты:\n"]
        for r in reversed(my_results):
            date, _, _, _, role_label, category, score, total, percent = r
            lines.append(f"{date} — {role_label} / {category}: {score}/{total} ({percent}%)")
        text = "\n".join(lines)

    if message_id:
        await bot.edit_message_text(text, chat_id=chat_id, message_id=message_id, reply_markup=back_to_menu_keyboard())
    else:
        await bot.send_message(chat_id, text)


@dp.callback_query(F.data == "menu:my_results")
async def menu_my_results(callback: types.CallbackQuery):
    await send_my_results(callback.message.chat.id, callback.from_user.id, message_id=callback.message.message_id)
    await callback.answer()


# =====================================================================
# ПРОХОЖДЕНИЕ ТЕСТА
# =====================================================================
@dp.callback_query(F.data.startswith("start_test:"))
async def start_test(callback: types.CallbackQuery):
    category = callback.data.split(":", 1)[1]
    role_key = get_user_role(callback.from_user.id)
    all_questions = load_questions_for_role(role_key)
    pool = [q for q in all_questions if q["category"] == category]
    random.shuffle(pool)
    count = get_questions_per_test(role_key)
    selected = pool if count == "all" else pool[:count]

    if not selected:
        await callback.answer("В этом разделе пока нет вопросов.", show_alert=True)
        return

    user_sessions[callback.from_user.id] = {
        "questions": selected,
        "index": 0,
        "score": 0,
        "category": category,
        "role_key": role_key,
        "chat_id": callback.message.chat.id,
        "message_id": callback.message.message_id,
    }
    await send_question(callback.from_user.id)
    await callback.answer()


async def send_question(user_id):
    session = user_sessions[user_id]
    index = session["index"]
    question = session["questions"][index]

    buttons = []
    for i, option in enumerate(question["options"], start=1):
        buttons.append([InlineKeyboardButton(text=str(option), callback_data=f"answer:{i}")])
    keyboard = InlineKeyboardMarkup(inline_keyboard=buttons)

    progress = "🟩" * index + "⬜️" * (len(session["questions"]) - index)
    text = f"{progress}\nВопрос {index + 1} из {len(session['questions'])}:\n\n{question['question']}"

    await bot.edit_message_text(
        text, chat_id=session["chat_id"], message_id=session["message_id"], reply_markup=keyboard
    )


@dp.callback_query(F.data.startswith("answer:"))
async def answer_handler(callback: types.CallbackQuery):
    user_id = callback.from_user.id
    if user_id not in user_sessions:
        await callback.answer("Тест не найден. Напиши /start, чтобы начать заново.")
        return

    session = user_sessions[user_id]
    chosen = int(callback.data.split(":", 1)[1])
    question = session["questions"][session["index"]]

    if chosen == question["correct"]:
        session["score"] += 1
        result_text = "✅ Верно!"
    else:
        correct_option = question["options"][question["correct"] - 1]
        result_text = f"❌ Неверно. Правильный ответ: {correct_option}"

    await callback.answer(result_text, show_alert=False)

    session["index"] += 1

    if session["index"] < len(session["questions"]):
        await send_question(user_id)
    else:
        score = session["score"]
        total = len(session["questions"])
        category = session["category"]
        role_key = session["role_key"]
        percent = round(score / total * 100)

        if percent >= 80:
            comment = "Отличный результат! 🎉"
        elif percent >= 50:
            comment = "Неплохо, но есть куда расти 💪"
        else:
            comment = "Стоит повторить материал 📖"

        text = (
            f"Тест завершён!\n\n"
            f"Раздел: {category}\n"
            f"Правильных ответов: {score} из {total} ({percent}%)\n\n"
            f"{comment}"
        )
        await bot.edit_message_text(
            text, chat_id=session["chat_id"], message_id=session["message_id"],
            reply_markup=back_to_menu_keyboard()
        )
        save_result(callback.from_user, role_key, category, score, total)
        del user_sessions[user_id]


# =====================================================================
# АДМИН-ПАНЕЛЬ
# =====================================================================
@dp.callback_query(F.data == "menu:admin")
async def menu_admin(callback: types.CallbackQuery):
    if not is_admin(callback.from_user.id):
        await callback.answer("Доступ только для администратора.", show_alert=True)
        return
    await callback.message.edit_text("⚙️ Админ-панель", reply_markup=admin_menu_keyboard())
    await callback.answer()


@dp.callback_query(F.data == "admin:stats")
async def admin_stats(callback: types.CallbackQuery):
    if not is_admin(callback.from_user.id):
        await callback.answer("Доступ только для администратора.", show_alert=True)
        return

    results = read_all_results()
    if not results:
        text = "Пока нет ни одного пройденного теста."
    else:
        total_tests = len(results)
        avg_percent = round(sum(r[8] for r in results) / total_tests)
        unique_users = len({r[1] for r in results})

        by_role = defaultdict(list)
        for r in results:
            by_role[r[4]].append(r[8])

        lines = [
            "📈 Общая статистика\n",
            f"Всего пройдено тестов: {total_tests}",
            f"Уникальных сотрудников: {unique_users}",
            f"Средний результат: {avg_percent}%\n",
            "По профессиям:",
        ]
        for role_label, percents in sorted(by_role.items()):
            avg = round(sum(percents) / len(percents))
            lines.append(f"  {role_label}: {avg}% (тестов: {len(percents)})")

        text = "\n".join(lines)

    await callback.message.edit_text(text, reply_markup=admin_menu_keyboard())
    await callback.answer()


@dp.callback_query(F.data == "admin:weak")
async def admin_weak(callback: types.CallbackQuery):
    if not is_admin(callback.from_user.id):
        await callback.answer("Доступ только для администратора.", show_alert=True)
        return

    results = read_all_results()
    weak = sorted([r for r in results if r[8] < 60], key=lambda r: r[8])[:15]

    if not weak:
        text = "Слабых результатов (ниже 60%) не найдено. 👍"
    else:
        lines = ["📉 Слабые результаты (ниже 60%):\n"]
        for r in weak:
            date, _, name, username, role_label, category, score, total, percent = r
            lines.append(f"{date} — {name} ({username}) — {role_label} / {category}: {score}/{total} ({percent}%)")
        text = "\n".join(lines)

    await callback.message.edit_text(text, reply_markup=admin_menu_keyboard())
    await callback.answer()


@dp.callback_query(F.data == "admin:download")
async def admin_download(callback: types.CallbackQuery):
    if not is_admin(callback.from_user.id):
        await callback.answer("Доступ только для администратора.", show_alert=True)
        return

    ensure_results_file()
    if os.path.exists(RESULTS_FILE):
        await bot.send_document(callback.message.chat.id, FSInputFile(RESULTS_FILE))
    await callback.answer()


@dp.callback_query(F.data == "admin:qcount")
async def admin_qcount(callback: types.CallbackQuery):
    if not is_admin(callback.from_user.id):
        await callback.answer("Доступ только для администратора.", show_alert=True)
        return
    await callback.message.edit_text(
        "Выбери профессию, чтобы изменить количество вопросов в тесте:",
        reply_markup=qcount_roles_keyboard()
    )
    await callback.answer()


@dp.callback_query(F.data.startswith("qcount_role:"))
async def qcount_role(callback: types.CallbackQuery):
    if not is_admin(callback.from_user.id):
        await callback.answer("Доступ только для администратора.", show_alert=True)
        return
    role_key = callback.data.split(":", 1)[1]
    label = ROLES[role_key]["label"]
    await callback.message.edit_text(
        f"Сколько вопросов задавать за один тест для «{label}»?",
        reply_markup=qcount_values_keyboard(role_key)
    )
    await callback.answer()


@dp.callback_query(F.data.startswith("qcount_set:"))
async def qcount_set(callback: types.CallbackQuery):
    if not is_admin(callback.from_user.id):
        await callback.answer("Доступ только для администратора.", show_alert=True)
        return
    _, role_key, value = callback.data.split(":", 2)
    value_to_save = value if value == "all" else int(value)
    save_questions_per_test(role_key, value_to_save)
    await callback.answer("Сохранено!")
    await callback.message.edit_text(
        "Выбери профессию, чтобы изменить количество вопросов в тесте:",
        reply_markup=qcount_roles_keyboard()
    )


# =====================================================================
# ЗАПУСК
# =====================================================================
async def main():
    await dp.start_polling(bot)


if __name__ == "__main__":
    asyncio.run(main())