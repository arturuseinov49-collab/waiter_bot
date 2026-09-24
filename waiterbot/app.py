"""Private-chat UI; durable training and reports share the same database."""
import asyncio
import logging
import weakref
from collections import defaultdict
from datetime import datetime, timedelta, timezone
from io import BytesIO

from aiogram import BaseMiddleware, Router, types
from aiogram.exceptions import TelegramBadRequest
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .catalog import ROLES, units
from .storage import LEGACY_RESULTS_HEADERS
from .training import Training

QUESTION_COUNTS = [5, 10, 15, 20, 25, 30]
MSK = timezone(timedelta(hours=3))


def keyboard(rows):
    return types.InlineKeyboardMarkup(inline_keyboard=[
        [types.InlineKeyboardButton(text=text, callback_data=data) for text, data in row]
        for row in rows
    ])


HOME = [[("⬅️ Главное меню", "menu:home")]]


def chunks(text, limit=3500):
    current, length = "", 0
    for char in text:
        width = units(char)
        if length + width > limit:
            yield current
            current, length = "", 0
        current += char
        length += width
    if current:
        yield current


def display_date(value):
    date = datetime.fromisoformat(value)
    if date.tzinfo is not None:
        date = date.astimezone(MSK)
    return date.strftime("%d.%m.%Y %H:%M")


def export_results(results):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Результаты"
    sheet.append(LEGACY_RESULTS_HEADERS + ["Режим"])
    for result in results:
        row = [display_date(result["created_at"]), result["user_id"], result["name"],
               "@" + result["username"] if result["username"] else "-",
               result["role_label"], result["category"], result["score"], result["total"],
               result["percent"], "Тренировка" if result["mode"] == "practice" else "Тест"]
        sheet.append(row)
        # User names and question categories must never become spreadsheet formulas.
        for cell in sheet[sheet.max_row]:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="254D73")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in zip("ABCDEFGHIJ", [20, 18, 30, 25, 22, 38, 14, 12, 12, 18]):
        sheet.column_dimensions[column].width = width
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


class Guard(BaseMiddleware):
    def __init__(self, app):
        self.app = app
        self.locks = weakref.WeakValueDictionary()

    async def __call__(self, handler, event, data):
        message = event.message if isinstance(event, types.CallbackQuery) else event
        user = event.from_user
        if not isinstance(message, types.Message) or not user:
            return
        if message.chat.type != "private":
            if isinstance(event, types.CallbackQuery):
                await event.answer("Открой личный чат с ботом.", show_alert=True)
            return
        if isinstance(event, types.CallbackQuery):
            try:
                await event.answer()
            except TelegramBadRequest:
                pass  # An old acknowledgement must not prevent recovery via /start.
        lock = self.locks.setdefault(user.id, asyncio.Lock())
        async with lock:
            try:
                await asyncio.to_thread(self.app.store.upsert_user, user.id, user.full_name, user.username)
                return await handler(event, data)
            except ValueError as error:
                await message.answer(str(error), reply_markup=keyboard(HOME))
            except Exception:
                logging.exception("Failed to handle a bot update")
                await message.answer(
                    "Не получилось выполнить действие. Нажми /start и «Продолжить».",
                    reply_markup=keyboard(HOME),
                )


class TrainingBot:
    def __init__(self, store, catalog, admin_id=0):
        self.store, self.catalog, self.admin_id = store, catalog, admin_id
        self.training = Training(store)
        self.router = Router()
        guard = Guard(self)
        self.router.message.outer_middleware(guard)
        self.router.callback_query.outer_middleware(guard)
        self.router.message.register(self.message)
        self.router.callback_query.register(self.callback)

    async def call(self, method, *args, **kwargs):
        return await asyncio.to_thread(method, *args, **kwargs)

    def is_admin(self, user_id):
        return bool(self.admin_id) and user_id == self.admin_id

    async def show(self, message, text, rows=None, edit=False):
        parts = list(chunks(text))
        markup = keyboard(rows or HOME)
        if edit and len(parts) == 1:
            try:
                await message.edit_text(text, reply_markup=markup)
                return
            except TelegramBadRequest as error:
                if "message is not modified" in error.message.lower():
                    return
                if not any(part in error.message.lower() for part in (
                    "message to edit not found", "message can't be edited", "message_id_invalid",
                )):
                    raise
        for index, part in enumerate(parts):
            await message.answer(part, reply_markup=markup if index == len(parts) - 1 else None)

    async def menu(self, message, user_id, edit=False):
        role = await self.call(self.store.get_role, user_id)
        session = await self.call(self.store.get_session, user_id)
        rows = []
        if session:
            rows += [[("▶️ Продолжить", "menu:resume")],
                     [("Завершить текущий тест", "cancel:" + session["id"])]]
        rows += [[("📝 Пройти тест", "menu:test")], [("🔄 Повторить ошибки", "menu:practice")],
                 [("📊 Мои результаты", "menu:my_results")], [("🔁 Сменить профессию", "menu:change_role")]]
        if self.is_admin(user_id):
            rows.append([("⚙️ Админ-панель", "menu:admin")])
        label = ROLES.get(role, {}).get("label", "не выбрана")
        text = f"Обучение команды\n\nТвоя профессия: {label}\n"
        if session:
            text += f"\nСохранён тест: {session['category']}\nОтвечено: {session['index']} из {len(session['questions'])}\n"
        text += "\nВ тесте можно сделать паузу: ответы и текущий вопрос сохраняются."
        await self.show(message, text, rows, edit)

    async def roles(self, message, user_id, edit=False):
        if await self.call(self.store.get_session, user_id):
            return await self.menu(message, user_id, edit)
        await self.show(message, "Выбери профессию:", [
            [(value["label"], "set_role:" + role)] for role, value in ROLES.items()
        ] + HOME, edit)

    async def categories(self, message, user_id, role=None, page=0, edit=False):
        selected = await self.call(self.store.get_role, user_id)
        if selected not in ROLES:
            return await self.roles(message, user_id, edit)
        if role is not None and role != selected:
            raise ValueError("Профессия изменилась. Открой список разделов заново.")
        if await self.call(self.store.get_session, user_id):
            return await self.menu(message, user_id, edit)
        items = list(self.catalog.categories(selected).items())
        pages = max(1, (len(items) + 7) // 8)
        page = max(0, min(page, pages - 1))
        rows = [[(category[:60], f"test:{selected}:{key}")] for key, category in items[page * 8:(page + 1) * 8]]
        nav = []
        if page:
            nav.append(("◀️", f"cats:{selected}:{page-1}"))
        if page + 1 < pages:
            nav.append(("▶️", f"cats:{selected}:{page+1}"))
        if nav:
            rows.append(nav)
        text = f"Выбери раздел для теста · {page+1}/{pages}" if items else "Вопросов для этой профессии пока нет."
        await self.show(message, text, rows + HOME, edit)

    async def start(self, message, user_id, role=None, category_id=None, practice=False, edit=False):
        selected = await self.call(self.store.get_role, user_id)
        if selected not in ROLES:
            return await self.roles(message, user_id, edit)
        if await self.call(self.store.get_session, user_id):
            return await self.menu(message, user_id, edit)
        if role is not None and role != selected:
            raise ValueError("Профессия изменилась. Открой список разделов заново.")
        if practice:
            saved = await self.call(self.store.get_mistakes, user_id, selected)
            questions = self.catalog.mistakes(selected, saved)
            category = "Повторение ошибок"
            if not questions:
                return await self.show(message, "Нет ошибок для повторения по текущей базе вопросов. Пройди тест, чтобы проверить знания.", HOME, edit)
        else:
            category, questions = self.catalog.select(selected, category_id)
        count = await self.call(self.store.get_question_count, selected)
        session = await self.call(self.training.start, user_id, selected, category, questions, count,
                                  "practice" if practice else "test")
        await self.render(message, user_id, session, edit)

    async def render(self, message, user_id, session, edit=False):
        index, total = session["index"], len(session["questions"])
        mode = "Тренировка" if session["mode"] == "practice" else "Тест"
        if session["phase"] == "complete":
            percent = round(session["score"] / total * 100)
            text = f"{mode} завершён!\n\n{session['category']}\nПравильных ответов: {session['score']} из {total} ({percent}%).\n\nРезультат сохранён."
            rows = [[("🔄 Повторить ошибки", "menu:practice")]] + HOME
        elif session["phase"] == "feedback":
            # Also completes a last answer saved just before an interrupted process.
            if index == total:
                await self.call(self.store.finish_session, user_id, session, keep_session=True)
            answer = session["answers"][-1]
            question = answer["question"]
            verdict = "✅ Верно!" if answer["correct"] else "❌ Есть ошибка"
            text = f"{verdict}\n\n{question['question']}\n\nТвой ответ: {question['options'][answer['chosen']-1]}\nПравильный ответ: {question['options'][question['correct']-1]}"
            if question.get("explanation"):
                text += "\n\n" + question["explanation"]
            rows = [[("Итоги" if index == total else "Следующий вопрос ➡️",
                      f"next:{session['id']}:{index}")]] + HOME
        else:
            question = session["questions"][index]
            progress = "".join("🟩" if a["correct"] else "🟥" for a in session["answers"][-20:])
            options = "\n\n".join(f"{number}. {option}" for number, option in enumerate(question["options"], 1))
            text = f"{mode} · {session['category']}\nВопрос {index+1} из {total}\n{progress}\n\n{question['question']}\n\n{options}"
            rows = [[(str(number), f"answer:{session['id']}:{index}:{number}") for number in range(1, 5)],
                    [("⏸ Пауза / меню", "menu:home")]]
        await self.show(message, text, rows, edit)

    async def results(self, message, user_id, edit=False):
        results = await self.call(self.store.read_results, user_id=user_id, limit=10)
        lines = ["📊 Последние результаты (время Москвы):"]
        for result in reversed(results):
            mode = "Тренировка" if result["mode"] == "practice" else "Тест"
            lines.append(f"{display_date(result['created_at'])} · {mode}\n{result['role_label']} / {result['category']}\n{result['score']}/{result['total']} ({result['percent']}%)")
        if not results:
            lines.append("Пока нет завершённых тестов.")
        await self.show(message, "\n\n".join(lines), HOME, edit)

    def admin_rows(self):
        return [[("📈 Общая статистика", "admin:stats")], [("📉 Слабые результаты", "admin:weak")],
                [("🔢 Количество вопросов", "admin:qcount")], [("📥 Скачать результаты", "admin:download")],
                [("🔄 Перечитать вопросы", "admin:reload")]] + HOME

    async def admin(self, message, user_id, action, edit=False):
        if not self.is_admin(user_id):
            raise ValueError("Доступ только для администратора.")
        rows = self.admin_rows()
        if action == "download":
            results = await self.call(self.store.read_results)
            contents = await self.call(export_results, results)
            await message.answer_document(types.BufferedInputFile(contents, filename="results.xlsx"))
            return
        if action == "reload":
            counts = await self.call(self.catalog.reload)
            text = "Вопросы проверены и обновлены:\n" + "\n".join(f"{ROLES[r]['label']}: {n}" for r, n in counts.items())
        elif action == "qcount":
            rows = []
            for role, details in ROLES.items():
                count = await self.call(self.store.get_question_count, role)
                rows.append([(f"{details['label']}: {'все' if count == 'all' else count}", "qcount_role:" + role)])
            rows += [[("⬅️ Админ-панель", "menu:admin")]]
            text = "Количество вопросов за один тест:"
        elif action in ("stats", "weak"):
            results = await self.call(self.store.read_results, mode="test")
            if action == "stats":
                by_role = defaultdict(list)
                for result in results:
                    by_role[result["role_label"]].append(result["percent"])
                average = round(sum(r["percent"] for r in results) / len(results)) if results else 0
                lines = ["📈 Статистика тестов", f"Тестов: {len(results)}",
                         f"Сотрудников: {len({r['user_id'] for r in results})}", f"Средний результат: {average}%",
                         "Тренировки по ошибкам считаются отдельно."]
                for role, scores in by_role.items():
                    lines.append(f"{role}: {round(sum(scores)/len(scores))}% · тестов {len(scores)}")
            else:
                # Recent results per employee/category show current needs, not old failures forever.
                latest = {(r["user_id"], r["role_key"], r["category"]): r for r in results}
                weak = sorted([r for r in latest.values() if r["percent"] < 60], key=lambda r: r["percent"])[:15]
                lines = ["📉 Последние тесты с результатом ниже 60%:"]
                lines += [f"{r['name']} · {r['role_label']} / {r['category']}: {r['percent']}% ({display_date(r['created_at'])})" for r in weak]
                if not weak:
                    lines.append("Таких результатов нет.")
            text = "\n\n".join(lines)
        else:
            text = "⚙️ Админ-панель"
        await self.show(message, text, rows, edit)

    async def message(self, message):
        user_id = message.from_user.id
        text = (message.text or "").split("@", 1)[0]
        if text in ("/test", "📝 Пройти тест"):
            return await self.categories(message, user_id)
        if text in ("/practice", "🔄 Повторить ошибки"):
            return await self.start(message, user_id, practice=True)
        if text in ("/results", "📊 Мои результаты"):
            return await self.results(message, user_id)
        if text in ("/admin", "⚙️ Админ-панель"):
            return await self.admin(message, user_id, "home")
        if text == "/resume":
            session = await self.call(self.store.get_session, user_id)
            if session:
                return await self.render(message, user_id, session)
        if text == "/start":
            await message.answer("Меню доступно снизу 👇", reply_markup=types.ReplyKeyboardMarkup(
                keyboard=[[types.KeyboardButton(text="📝 Пройти тест"), types.KeyboardButton(text="🔄 Повторить ошибки")],
                          [types.KeyboardButton(text="📊 Мои результаты"), types.KeyboardButton(text="🏠 Меню")]],
                resize_keyboard=True, is_persistent=True))
        if await self.call(self.store.get_role, user_id) not in ROLES:
            return await self.roles(message, user_id)
        await self.menu(message, user_id)

    async def callback(self, callback):
        message, user_id = callback.message, callback.from_user.id
        parts = (callback.data or "").split(":")
        prefix = parts[0]
        if prefix in ("admin", "qcount_role", "qcount_set") and not self.is_admin(user_id):
            raise ValueError("Доступ только для администратора.")
        if prefix == "menu" and len(parts) == 2:
            action = parts[1]
            if action == "home":
                return await self.menu(message, user_id, True)
            if action == "change_role":
                return await self.roles(message, user_id, True)
            if action == "test":
                return await self.categories(message, user_id, edit=True)
            if action == "practice":
                return await self.start(message, user_id, practice=True, edit=True)
            if action == "my_results":
                return await self.results(message, user_id, True)
            if action == "admin":
                return await self.admin(message, user_id, "home", True)
            if action == "resume":
                session = await self.call(self.store.get_session, user_id)
                if session:
                    return await self.render(message, user_id, session, True)
                return await self.menu(message, user_id, True)
        if prefix == "set_role" and len(parts) == 2 and parts[1] in ROLES:
            if await self.call(self.store.get_session, user_id):
                return await self.menu(message, user_id, True)
            await self.call(self.store.set_role, user_id, parts[1])
            return await self.menu(message, user_id, True)
        if prefix == "cats" and len(parts) == 3:
            return await self.categories(message, user_id, parts[1], int(parts[2]), True)
        if prefix == "test" and len(parts) == 3:
            return await self.start(message, user_id, parts[1], parts[2], edit=True)
        if prefix == "answer" and len(parts) == 4:
            session = await self.call(self.training.answer, user_id, parts[1], int(parts[2]), int(parts[3]))
            return await self.render(message, user_id, session, True)
        if prefix == "next" and len(parts) == 3:
            session = await self.call(self.training.next, user_id, parts[1], int(parts[2]))
            return await self.render(message, user_id, session, True)
        if prefix == "cancel" and len(parts) == 2:
            session = await self.call(self.store.get_session, user_id)
            if not session or session["id"] != parts[1]:
                raise ValueError("Этот тест уже закрыт.")
            return await self.show(message, "Закрыть текущий тест? Если он не закончен, его прогресс будет удалён.",
                                   [[("Да, закрыть", "cancel_yes:" + parts[1])], [("Продолжить", "menu:resume")]] + HOME, True)
        if prefix == "cancel_yes" and len(parts) == 2:
            await self.call(self.training.cancel, user_id, parts[1])
            return await self.menu(message, user_id, True)
        if prefix == "admin" and len(parts) == 2:
            return await self.admin(message, user_id, parts[1], True)
        if prefix == "qcount_role" and len(parts) == 2 and parts[1] in ROLES:
            rows = [[(str(n), f"qcount_set:{parts[1]}:{n}") for n in QUESTION_COUNTS[:3]],
                    [(str(n), f"qcount_set:{parts[1]}:{n}") for n in QUESTION_COUNTS[3:]],
                    [("Все вопросы", f"qcount_set:{parts[1]}:all")], [("⬅️ Назад", "admin:qcount")]]
            return await self.show(message, "Сколько вопросов включать в тест?", rows, True)
        if prefix == "qcount_set" and len(parts) == 3 and parts[1] in ROLES and parts[2] in {str(n) for n in QUESTION_COUNTS} | {"all"}:
            await self.call(self.store.set_question_count, parts[1], parts[2])
            return await self.admin(message, user_id, "qcount", True)
        raise ValueError("Эта кнопка больше не действует. Открой /start.")
