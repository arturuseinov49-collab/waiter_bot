"""Validated, versioned question snapshots from the restaurant's workbooks."""
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook

ROLES = {
    "waiter": {"label": "🍽 Официант", "file": "questions_waiter.xlsx"},
    "pizzaiolo": {"label": "🍕 Пиццайоло", "file": "questions_pizzaiolo.xlsx"},
    "cook": {"label": "👨‍🍳 Повар", "file": "questions_cook.xlsx"},
}
HEADERS = ["Категория", "Вопрос", "Вариант 1", "Вариант 2", "Вариант 3", "Вариант 4", "Номер правильного (1-4)"]


def identity(value):
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True).encode()).hexdigest()[:16]


def units(text):
    return len(text.encode("utf-16-le")) // 2


class Catalog:
    def __init__(self, base_dir):
        self.base_dir = Path(base_dir)
        self.questions = {}

    def reload(self):
        replacement = {role: self._read(role) for role in ROLES}
        self.questions = replacement
        return {role: len(items) for role, items in replacement.items()}

    def _read(self, role):
        path = self.base_dir / ROLES[role]["file"]
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            sheet = workbook.active
            rows = sheet.iter_rows(values_only=True)
            headers = list(next(rows, ()))
            if headers[:7] != HEADERS:
                raise ValueError(f"{path.name}: неверные заголовки семи основных столбцов")
            explanation_index = headers.index("Объяснение") if "Объяснение" in headers else None
            questions, seen = [], set()
            for number, row in enumerate(rows, 2):
                if all(value is None for value in row):
                    continue
                prefix = f"{path.name}, строка {number}"
                values = row[:6]
                if len(row) < 7 or any(not isinstance(v, str) or not v.strip() or v.startswith("=") for v in values):
                    raise ValueError(f"{prefix}: категория, вопрос и четыре ответа должны содержать текст без формул")
                category, question, *options = [v.strip() for v in values]
                correct = row[6]
                if isinstance(correct, bool) or str(correct).strip() not in {"1", "2", "3", "4", "1.0", "2.0", "3.0", "4.0"}:
                    raise ValueError(f"{prefix}: правильный ответ должен быть числом от 1 до 4")
                if len(set(o.casefold() for o in options)) != 4:
                    raise ValueError(f"{prefix}: повторяющиеся варианты ответа")
                explanation = row[explanation_index] if explanation_index is not None else None
                if explanation is not None and (not isinstance(explanation, str) or explanation.startswith("=")):
                    raise ValueError(f"{prefix}: объяснение должно быть текстом")
                explanation = (explanation or "").strip()
                if units(question + "\n".join(options) + explanation) > 3100 or units(category) > 150:
                    raise ValueError(f"{prefix}: вопрос или объяснение слишком длинные для Telegram")
                item = {"category": category, "question": question, "options": options,
                        "correct": int(float(correct)), "explanation": explanation}
                item["id"] = identity({"role": role, **item})
                if item["id"] not in seen:
                    questions.append(item)
                    seen.add(item["id"])
            return questions
        finally:
            workbook.close()

    def categories(self, role):
        return {identity(category): category for category in sorted({q["category"] for q in self.questions[role]})}

    def select(self, role, category_id):
        category = self.categories(role).get(category_id)
        if category is None:
            raise ValueError("Раздел изменился. Открой список разделов заново.")
        return category, [q for q in self.questions[role] if q["category"] == category]

    def mistakes(self, role, saved):
        current = {q["id"]: q for q in self.questions[role]}
        return [current[q["id"]] for q in saved if q["id"] in current]
