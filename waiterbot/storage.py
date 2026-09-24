"""Durable state for the restaurant bot.

Each public operation owns its connection, so callers may safely run the
synchronous API with ``asyncio.to_thread``. Finalizing a test is one transaction.
"""

from __future__ import annotations

import hashlib
import json
import os
import sqlite3
import tempfile
from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterator


LEGACY_RESULTS_HEADERS = [
    "Дата и время", "User ID", "Имя", "Username", "Профессия", "Раздел",
    "Правильных", "Всего", "Процент",
]


def _integer(value: Any, field: str, minimum: int = 0) -> int:
    if isinstance(value, bool):
        raise ValueError(f"{field}: требуется целое число")
    if isinstance(value, int):
        result = value
    elif isinstance(value, float) and value.is_integer():
        result = int(value)
    elif isinstance(value, str) and value.strip().isdigit():
        result = int(value.strip())
    else:
        raise ValueError(f"{field}: требуется целое число")
    if result < minimum:
        raise ValueError(f"{field}: значение должно быть не меньше {minimum}")
    return result


def _text(value: Any, field: str) -> str:
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"{field}: требуется непустая строка")
    return value


def _count(value: Any) -> int | str:
    return "all" if value == "all" else _integer(value, "Количество вопросов", 1)


def _json(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, allow_nan=False, separators=(",", ":"))


def _now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def _legacy_date(value: Any) -> str:
    if isinstance(value, datetime):
        return value.isoformat(timespec="seconds")
    value = _text(value, "Дата результата")
    try:
        return datetime.strptime(value, "%d.%m.%Y %H:%M").isoformat(timespec="seconds")
    except ValueError:
        try:
            return datetime.fromisoformat(value).isoformat(timespec="seconds")
        except ValueError as exc:
            raise ValueError("Неизвестный формат даты результата") from exc


class Store:
    def __init__(self, db_path: Path):
        self.db_path = Path(db_path).resolve()

    @contextmanager
    def _connection(self, *, write: bool = False) -> Iterator[sqlite3.Connection]:
        connection = sqlite3.connect(self.db_path, timeout=30)
        connection.row_factory = sqlite3.Row
        try:
            connection.execute("PRAGMA busy_timeout = 30000")
            connection.execute("PRAGMA foreign_keys = ON")
            if write:
                connection.execute("BEGIN IMMEDIATE")
            yield connection
            connection.commit()
        except BaseException:
            connection.rollback()
            raise
        finally:
            connection.close()

    def initialize(self) -> None:
        self.db_path.parent.mkdir(parents=True, exist_ok=True)
        with self._connection() as connection:
            connection.execute("PRAGMA journal_mode = WAL")
            connection.executescript("""
                CREATE TABLE IF NOT EXISTS users (
                    user_id INTEGER PRIMARY KEY,
                    name TEXT NOT NULL DEFAULT '',
                    username TEXT,
                    role_key TEXT
                );
                CREATE TABLE IF NOT EXISTS role_labels (
                    role_key TEXT PRIMARY KEY,
                    label TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS settings (
                    role_key TEXT PRIMARY KEY,
                    question_count TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS sessions (
                    user_id INTEGER PRIMARY KEY REFERENCES users(user_id),
                    session_id TEXT NOT NULL,
                    payload TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS results (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    session_id TEXT NOT NULL UNIQUE,
                    created_at TEXT NOT NULL,
                    user_id INTEGER NOT NULL REFERENCES users(user_id),
                    name TEXT NOT NULL,
                    username TEXT,
                    role_key TEXT NOT NULL,
                    role_label TEXT NOT NULL,
                    category TEXT NOT NULL,
                    score INTEGER NOT NULL CHECK (score >= 0),
                    total INTEGER NOT NULL CHECK (total > 0 AND score <= total),
                    percent INTEGER NOT NULL CHECK (percent BETWEEN 0 AND 100),
                    mode TEXT NOT NULL CHECK (mode IN ('test', 'practice')),
                    answers_json TEXT NOT NULL DEFAULT '[]'
                );
                CREATE INDEX IF NOT EXISTS results_user_mode ON results(user_id, mode, created_at);
                CREATE TABLE IF NOT EXISTS mistakes (
                    user_id INTEGER NOT NULL REFERENCES users(user_id),
                    role_key TEXT NOT NULL,
                    question_id TEXT NOT NULL,
                    question_json TEXT NOT NULL,
                    misses INTEGER NOT NULL DEFAULT 0,
                    resolved INTEGER NOT NULL DEFAULT 0 CHECK (resolved IN (0, 1)),
                    updated_at TEXT NOT NULL,
                    PRIMARY KEY (user_id, role_key, question_id)
                );
                CREATE TABLE IF NOT EXISTS metadata (
                    key TEXT PRIMARY KEY,
                    value TEXT NOT NULL
                );
            """)

    @staticmethod
    def _ensure_user(connection: sqlite3.Connection, user_id: int) -> None:
        connection.execute("INSERT OR IGNORE INTO users(user_id) VALUES (?)", (user_id,))

    def upsert_user(self, user_id: int, name: str, username: str | None) -> None:
        user_id = _integer(user_id, "User ID", 1)
        if not isinstance(name, str) or (username is not None and not isinstance(username, str)):
            raise ValueError("Имя и username должны быть строками")
        with self._connection(write=True) as connection:
            connection.execute("""
                INSERT INTO users(user_id, name, username) VALUES (?, ?, ?)
                ON CONFLICT(user_id) DO UPDATE SET name=excluded.name, username=excluded.username
            """, (user_id, name, username))

    def get_role(self, user_id: int) -> str | None:
        with self._connection() as connection:
            row = connection.execute("SELECT role_key FROM users WHERE user_id=?", (user_id,)).fetchone()
            return row[0] if row else None

    def set_role(self, user_id: int, role: str) -> None:
        user_id = _integer(user_id, "User ID", 1)
        role = _text(role, "Профессия")
        with self._connection(write=True) as connection:
            self._ensure_user(connection, user_id)
            connection.execute("UPDATE users SET role_key=? WHERE user_id=?", (role, user_id))

    def get_question_count(self, role: str, default: int = 10) -> int | str:
        with self._connection() as connection:
            row = connection.execute("SELECT question_count FROM settings WHERE role_key=?", (role,)).fetchone()
            return _count(row[0]) if row else _count(default)

    def set_question_count(self, role: str, value: int | str) -> None:
        role, value = _text(role, "Профессия"), _count(value)
        with self._connection(write=True) as connection:
            connection.execute("""
                INSERT INTO settings(role_key, question_count) VALUES (?, ?)
                ON CONFLICT(role_key) DO UPDATE SET question_count=excluded.question_count
            """, (role, str(value)))

    def get_session(self, user_id: int) -> dict | None:
        with self._connection() as connection:
            row = connection.execute("SELECT payload FROM sessions WHERE user_id=?", (user_id,)).fetchone()
            return json.loads(row[0]) if row else None

    def save_session(self, user_id: int, session: dict) -> None:
        user_id = _integer(user_id, "User ID", 1)
        session_id = _text(session.get("id"), "ID сессии")
        payload = _json(session)
        with self._connection(write=True) as connection:
            self._ensure_user(connection, user_id)
            connection.execute("""
                INSERT INTO sessions(user_id, session_id, payload) VALUES (?, ?, ?)
                ON CONFLICT(user_id) DO UPDATE SET session_id=excluded.session_id, payload=excluded.payload
            """, (user_id, session_id, payload))

    def delete_session(self, user_id: int) -> None:
        with self._connection(write=True) as connection:
            connection.execute("DELETE FROM sessions WHERE user_id=?", (user_id,))


    def compare_session(self, user_id: int, expected: dict | None, replacement: dict | None) -> bool:
        """Atomically save an action only if the persisted snapshot is still current."""
        user_id = _integer(user_id, "User ID", 1)
        with self._connection(write=True) as connection:
            row = connection.execute("SELECT payload FROM sessions WHERE user_id=?", (user_id,)).fetchone()
            current = json.loads(row[0]) if row else None
            if current != expected:
                return False
            self._ensure_user(connection, user_id)
            if replacement is None:
                connection.execute("DELETE FROM sessions WHERE user_id=?", (user_id,))
            else:
                session_id = _text(replacement.get("id"), "ID сессии")
                connection.execute("""INSERT INTO sessions(user_id,session_id,payload) VALUES (?,?,?)
                    ON CONFLICT(user_id) DO UPDATE SET session_id=excluded.session_id,payload=excluded.payload""",
                    (user_id, session_id, _json(replacement)))
            return True

    @staticmethod
    def _completed_session(session: dict) -> tuple[str, str, str, str, int, int, list[dict]]:
        session_id = _text(session.get("id"), "ID сессии")
        role = _text(session.get("role_key"), "Профессия")
        category = _text(session.get("category"), "Раздел")
        mode = session.get("mode", "test")
        if mode not in ("test", "practice"):
            raise ValueError("Неизвестный режим прохождения")
        questions, answers = session.get("questions"), session.get("answers")
        if not isinstance(questions, list) or not questions or not isinstance(answers, list):
            raise ValueError("Отсутствуют вопросы или ответы завершённого теста")
        total = len(questions)
        if len(answers) != total or session.get("index") != total:
            raise ValueError("Нельзя сохранить незавершённый тест как результат")
        for question, answer in zip(questions, answers):
            if not isinstance(question, dict) or not isinstance(answer, dict) or answer.get("question") != question:
                raise ValueError("Ответ не соответствует вопросу сессии")
            _text(question.get("id"), "ID вопроса")
            _text(question.get("question"), "Текст вопроса")
            options = question.get("options")
            if not isinstance(options, list) or len(options) < 2 or any(not isinstance(o, str) or not o.strip() for o in options):
                raise ValueError("Некорректные варианты ответа")
            correct = _integer(question.get("correct"), "Правильный ответ", 1)
            chosen = _integer(answer.get("chosen"), "Выбранный ответ", 1)
            if correct > len(options) or chosen > len(options):
                raise ValueError("Индекс ответа выходит за границы вариантов")
            if not isinstance(answer.get("correct"), bool) or answer["correct"] != (chosen == correct):
                raise ValueError("Некорректная оценка ответа")
        score = sum(answer["correct"] for answer in answers)
        if _integer(session.get("score"), "Результат") != score:
            raise ValueError("Результат не соответствует ответам")
        return session_id, role, category, mode, score, total, answers

    def finish_session(self, user_id: int, session: dict, *, keep_session: bool = False) -> bool:
        """Record one completed attempt; return False for an already saved ID.

        A stale completion never deletes a different, newer active session.
        Both test and practice answers update the personal mistake bank.
        """
        user_id = _integer(user_id, "User ID", 1)
        session_id, role, category, mode, score, total, answers = self._completed_session(session)
        with self._connection(write=True) as connection:
            self._ensure_user(connection, user_id)
            user = connection.execute("SELECT name, username FROM users WHERE user_id=?", (user_id,)).fetchone()
            label = connection.execute("SELECT label FROM role_labels WHERE role_key=?", (role,)).fetchone()
            created_at = _now()
            cursor = connection.execute("""
                INSERT INTO results(session_id, created_at, user_id, name, username, role_key,
                                    role_label, category, score, total, percent, mode, answers_json)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                ON CONFLICT(session_id) DO NOTHING
            """, (session_id, created_at, user_id, user["name"], user["username"], role,
                  label[0] if label else role, category, score, total, round(score / total * 100), mode, _json(answers)))
            saved = cursor.rowcount == 1
            if saved:
                for answer in answers:
                    question = answer["question"]
                    if answer["correct"]:
                        connection.execute("""
                            UPDATE mistakes SET resolved=1, updated_at=?, question_json=?
                            WHERE user_id=? AND role_key=? AND question_id=?
                        """, (created_at, _json(question), user_id, role, question["id"]))
                    else:
                        connection.execute("""
                            INSERT INTO mistakes(user_id, role_key, question_id, question_json, misses, resolved, updated_at)
                            VALUES (?, ?, ?, ?, 1, 0, ?)
                            ON CONFLICT(user_id, role_key, question_id) DO UPDATE SET
                                question_json=excluded.question_json, misses=mistakes.misses+1,
                                resolved=0, updated_at=excluded.updated_at
                        """, (user_id, role, question["id"], _json(question), created_at))
            if not keep_session:
                connection.execute("DELETE FROM sessions WHERE user_id=? AND session_id=?", (user_id, session_id))
            return saved

    def read_results(self, user_id: int | None = None, mode: str | None = None, limit: int | None = None) -> list[dict]:
        """Return the latest ``limit`` matching results in chronological order.

        Use mode='test' for official assessment statistics; None includes practice.
        """
        clauses, parameters = [], []
        if user_id is not None:
            clauses.append("user_id=?")
            parameters.append(user_id)
        if mode is not None:
            if mode not in ("test", "practice"):
                raise ValueError("Неизвестный режим прохождения")
            clauses.append("mode=?")
            parameters.append(mode)
        query = """SELECT id,session_id,created_at,user_id,name,username,role_key,role_label,
                          category,score,total,percent,mode FROM results"""
        if clauses:
            query += " WHERE " + " AND ".join(clauses)
        query += " ORDER BY julianday(created_at) DESC, id DESC"
        if limit is not None:
            query += " LIMIT ?"
            parameters.append(_integer(limit, "Лимит результатов"))
        with self._connection() as connection:
            return [dict(row) for row in reversed(connection.execute(query, parameters).fetchall())]

    def get_mistakes(self, user_id: int, role_key: str | None = None) -> list[dict]:
        query = "SELECT question_json FROM mistakes WHERE user_id=? AND resolved=0"
        parameters: list[Any] = [user_id]
        if role_key is not None:
            query += " AND role_key=?"
            parameters.append(role_key)
        query += " ORDER BY misses DESC, updated_at DESC, question_id"
        with self._connection() as connection:
            return [json.loads(row[0]) for row in connection.execute(query, parameters)]

    def migrate_legacy(self, base_dir: Path, roles: dict) -> dict:
        """Import existing legacy files once, without changing or removing them.

        Each source is marked only after a successful, atomic transaction. Unknown
        headers, malformed rows or invalid JSON abort all imports in this call.
        Missing sources stay eligible for import if supplied on a later startup.
        """
        from openpyxl import load_workbook

        base_dir = Path(base_dir)
        labels = {key: _text(value["label"], "Название профессии") for key, value in roles.items()}
        by_label = {value: key for key, value in labels.items()}
        counts = {"roles": 0, "settings": 0, "results": 0}
        with self._connection(write=True) as connection:
            for key, label in labels.items():
                connection.execute("""INSERT INTO role_labels(role_key,label) VALUES (?,?)
                    ON CONFLICT(role_key) DO UPDATE SET label=excluded.label""", (key, label))
            for filename, count_key in (("user_roles.json", "roles"), ("settings.json", "settings"), ("results.xlsx", "results")):
                path = base_dir / filename
                marker = f"legacy_import_v1:{filename}"
                if not path.exists() or connection.execute("SELECT 1 FROM metadata WHERE key=?", (marker,)).fetchone():
                    continue
                try:
                    digest = hashlib.sha256(path.read_bytes()).hexdigest()
                    if filename.endswith(".json"):
                        data = json.loads(path.read_text(encoding="utf-8-sig"))
                        if not isinstance(data, dict):
                            raise ValueError("Корневое значение JSON должно быть объектом")
                        for key, value in data.items():
                            if filename == "user_roles.json":
                                user_id = _integer(key, "User ID", 1)
                                if value not in roles:
                                    raise ValueError("Неизвестная профессия в user_roles.json")
                                self._ensure_user(connection, user_id)
                                connection.execute("UPDATE users SET role_key=COALESCE(role_key, ?) WHERE user_id=?", (value, user_id))
                            else:
                                if key not in roles:
                                    raise ValueError("Неизвестная профессия в settings.json")
                                connection.execute("INSERT OR IGNORE INTO settings(role_key,question_count) VALUES (?,?)", (key, str(_count(value))))
                            counts[count_key] += 1
                    else:
                        workbook = load_workbook(path, read_only=True, data_only=False)
                        try:
                            sheet = workbook.active
                            if sheet is None:
                                raise ValueError("В файле результатов нет листа")
                            rows = sheet.iter_rows(values_only=True)
                            if list(next(rows, ())) != LEGACY_RESULTS_HEADERS:
                                raise ValueError("Неизвестные заголовки results.xlsx; требуется ручная миграция без потери данных")
                            for number, row in enumerate(rows, start=2):
                                if all(value is None for value in row):
                                    continue
                                if len(row) != len(LEGACY_RESULTS_HEADERS):
                                    raise ValueError(f"Строка {number}: неверное количество столбцов")
                                date, user_id, name, username, label, category, score, total, percent = row
                                user_id = _integer(user_id, "User ID", 1)
                                name, label, category = _text(name, "Имя"), _text(label, "Профессия"), _text(category, "Раздел")
                                if label not in by_label and label not in roles:
                                    raise ValueError(f"Строка {number}: неизвестная профессия")
                                role = by_label.get(label, label)
                                score, total, percent = _integer(score, "Результат"), _integer(total, "Всего", 1), _integer(percent, "Процент")
                                if score > total or percent > 100 or percent != round(score / total * 100):
                                    raise ValueError(f"Строка {number}: некорректный результат")
                                if username in (None, "", "-"):
                                    username = None
                                else:
                                    username = _text(username, "Username").lstrip("@")
                                self._ensure_user(connection, user_id)
                                connection.execute("""UPDATE users SET
                                    name=CASE WHEN name='' THEN ? ELSE name END,
                                    username=COALESCE(username, ?) WHERE user_id=?""", (name, username, user_id))
                                connection.execute("""
                                    INSERT INTO results(session_id,created_at,user_id,name,username,role_key,
                                        role_label,category,score,total,percent,mode)
                                    VALUES (?,?,?,?,?,?,?,?,?,?,?,'test')
                                """, (f"legacy:{digest}:{number}", _legacy_date(date), user_id, name, username,
                                      role, label, category, score, total, percent))
                                counts["results"] += 1
                        finally:
                            workbook.close()
                    connection.execute("INSERT INTO metadata(key,value) VALUES (?,?)", (marker, _json({"sha256": digest, "imported_at": _now()})))
                except Exception as exc:
                    raise ValueError(f"Не удалось импортировать {filename}: {exc}") from exc
        return counts

    def backup(self, destination: Path) -> None:
        """Create a consistent SQLite snapshot and atomically publish it."""
        destination = Path(destination).resolve()
        if destination == self.db_path:
            raise ValueError("Резервная копия не может заменить рабочую базу")
        destination.parent.mkdir(parents=True, exist_ok=True)
        descriptor, temporary = tempfile.mkstemp(prefix=f".{destination.name}.", suffix=".tmp", dir=destination.parent)
        os.close(descriptor)
        try:
            with self._connection() as source:
                target = sqlite3.connect(temporary)
                try:
                    source.backup(target)
                finally:
                    target.close()
            os.replace(temporary, destination)
        finally:
            if os.path.exists(temporary):
                os.unlink(temporary)
